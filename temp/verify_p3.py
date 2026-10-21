# Table-aware verification of P3-EVHM_paper.docx: structure, authors, sections, equations, refs, two-column layout, figure, tables.
import os, sys
from docx import Document
from docx.oxml.ns import qn

BASE = r"C:\Users\jithe\p3-evhm-paper"
path = os.path.join(BASE, "P3-EVHM_paper.docx")
doc = Document(path)

def all_text(d):
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)

txt = all_text(doc)
checks, fails = 0, []

def chk(name, cond):
    global checks
    checks += 1
    if not cond: fails.append(name)
    print(("[OK]   " if cond else "[FAIL] ") + name)

# Two-column section check
two_col = 0
for s in doc.sections:
    cols = s._sectPr.xpath("./w:cols")
    if cols and cols[0].get(qn("w:num")) == "2": two_col += 1
chk("two-column body section present", two_col >= 1)

# Images
img_count = len(doc.inline_shapes)
chk("architecture figure embedded", img_count >= 1)
chk("results figure embedded (2 figures total)", img_count >= 2)

# Title/authors/affiliation
chk("title present", "Parallel Physics-Informed Prognostics for EV Battery Health Management (P3-EVHM)" in txt)
chk("subtitle present", "Integrating Physics-Informed LSTM Networks" in txt)
chk("author A. Vibhas Chowdhary", "A. Vibhas Chowdhary" in txt)
chk("author Maddala Jithendra", "Maddala Jithendra" in txt)
chk("guide Kavitha R", "Kavitha R" in txt)
chk("SRM affiliation", "SRM Institute of Science and Technology" in txt)
chk("email va1641", "va1641@srmist.edu.in" in txt)
chk("email jm4408", "jm4408@srmist.edu.in" in txt)
chk("email kavithar14", "kavithar14@srmist.edu.in" in txt)

# Sections
for h in ["I. INTRODUCTION", "II. RELATED WORK AND RESEARCH GAP",
          "III. DATASET AND PRELIMINARY DATA ANALYSIS", "IV. SYSTEM ARCHITECTURE",
          "V. PROPOSED METHODOLOGY", "VI. GPU PARALLEL IMPLEMENTATION",
          "VII. BENCHMARK AND EVALUATION PLAN", "VIII. VALIDATION RESULTS (CPU AND GPU)",
          "IX. FUNCTIONAL AND NON-FUNCTIONAL REQUIREMENTS",
          "X. EXPECTED OUTCOMES AND LIMITATIONS", "XI. CONCLUSION",
          "REFERENCES", "ACKNOWLEDGMENT"]:
    chk("section: " + h, h in txt)

# Subsections
for h in ["A. Physics-informed deep learning for batteries", "B. Particle-filter-based battery prognostics",
          "C. GPU-accelerated prognostics", "D. Agentic and distributed PHM", "E. Research gap",
          "A. Problem setup", "B. Physics-informed LSTM", "C. Particle-filter RUL estimation",
          "D. Physics-consistency metrics"]:
    chk("subsection: " + h, h in txt)

# Abstract, keywords, honest framing
chk("Abstract block", "Abstract-" in txt)
chk("Index Terms block", "Index Terms-" in txt)
chk("honest: all results measured", "All reported results are measured on the NASA PCoE dataset" in txt)
chk("gap statement present", "no published framework integrates physics-informed LSTM training" in txt)
chk("TL-MLSTM-PF future-work quote", "incorporating physics-informed constraints" in txt)
chk("naming collision note (PI-LSTM)", "PI-LSTM is already used" in txt)
chk("naming collision note (parallel topology)", "denotes network branch topology" in txt)
chk("naming collision note (SP-LSTM)", "single-particle electrochemical model" in txt)

# Equations (1)-(8)
for n in ["(1)", "(2)", "(3)", "(4)", "(5)", "(6)", "(7)", "(8)"]:
    chk("equation " + n + " present", n in txt)

# Tables
chk("Table I evaluation metrics", "TABLE I. Evaluation metrics" in txt)
chk("Table II budget", "TABLE II. Indicative budget" in txt)
chk("Table III grid summary", "TABLE III. Grid summary" in txt)
chk("Table IV GPU scaling", "TABLE IV. Measured GPU scaling" in txt)
chk("budget total row", "3,000-10,500" in txt)
chk("Table I has coverage metric", "95% interval coverage / width" in txt)
chk("Table I has violation rate", "Constraint-violation rate" in txt)
chk("Table I has speedup", "Speedup" in txt)
chk("grid: T4 mean V B2", "0.366" in txt)
chk("grid: T4 mean V B3", "0.076" in txt)
chk("grid: V reduced 15 of 20", "15 of 20" in txt)
chk("grid: B5 18.6 (E1)", "18.6 cycles (20/20)" in txt)
chk("grid: B5 18.8 (E2)", "18.8 cycles (20/20)" in txt)
chk("grid: B5 coverage 0.85 both", "0.85" in txt)
chk("grid: B6 weakness 5/20 and 10/20", "32.8 cycles (5/20)" in txt and "30.9 cycles (10/20)" in txt)
chk("GPU scaling 16.5x at 500k", "16.5x" in txt)
chk("GPU parity 0.98x at 5k", "0.98x" in txt)
chk("GPU small-N honesty (warmup note)", "warmup" in txt)
chk("kernel URL present", "kaggle.com/code/jithendramaddala/p3-evhm-gpu-run" in txt)
chk("cross-cell finding", "Cross-cell generalization" in txt)
chk("finding: 3-seed repeatability", "3-seed repeatability" in txt)
chk("Fig. 2 caption", "Fig. 2. Measured trajectories" in txt)

chk("transfer: pretrained mode subsection", "E. Pretrained transfer mode" in txt)
chk("transfer: 12 transfer runs RUL", "all 12 transfer runs" in txt)
chk("transfer: scratch RUL only 1 of 4", "only 1 of 4 held-out cells" in txt)
chk("transfer: Table V present", "TABLE V. Leave-one-cell-out transfer validation" in txt)
chk("transfer: frozen body", "freezes every LSTM body layer" in txt)

# References
import re
ref_nums = re.findall(r"^\[(\d+|26b)\]", txt, re.M)
chk("reference entries >= 30", len(ref_nums) >= 30)
for key, label in [("[13]", "Shi PI-LSTM MSSP"), ("[14]", "TL-MLSTM-PF Batteries"),
                   ("[15]", "SPF Measurement"), ("[16]", "GPU Accelerated Prognostics PHM 2014"),
                   ("[17]", "CRPF CUDA"), ("[18]", "Metropolis resampler"),
                   ("[22]", "hybrid Bayesian PINN"), ("[25]", "PHMForge"),
                   ("[29]", "Dendrite CUDA"), ("[32]", "SP-LSTM single particle")]:
    chk("ref " + key + " " + label, key in txt)
chk("ref [13] has MSSP volume", "109347" in txt)
chk("ref [13] has URL", "S0888327022004824" in txt)
chk("ref [16] has PHM proceedings URL", "papers.phmsociety.org" in txt)
chk("ref [20] has Sci Rep DOI", "10.1038/s41598-026-37850-y" in txt)
chk("incomplete refs flagged honestly", "Incomplete" in txt)

# Word count stats
words = len(txt.split())
print("WORDS", words, "| PARAGRAPHS", len(doc.paragraphs), "| TABLES", len(doc.tables),
      "| IMAGES", img_count, "| SIZE", os.path.getsize(path), "bytes")
print("RESULT:", "PASS " + str(checks - len(fails)) + "/" + str(checks) + " checks, " + str(len(fails)) + " failures")
if fails:
    print("FAILED:", "; ".join(fails)); sys.exit(1)
